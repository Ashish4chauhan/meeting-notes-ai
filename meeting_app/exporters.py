import os
from pathlib import Path


def export_meeting(path: str | Path, title: str, transcript: str, summary: str, segments=None):
    destination = Path(path).expanduser()
    suffix = destination.suffix.lower()
    content = f"{title}\n{'=' * len(title)}\n\n{summary}\n\nTranscript\n----------\n{transcript}\n"
    if suffix == ".txt":
        destination.write_text(content, encoding="utf-8")
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX export requires python-docx. Install dependencies from requirements.txt.") from exc
        doc = Document()
        doc.add_heading(title, 0)
        doc.add_heading("Summary", level=1)
        for line in summary.splitlines():
            if line.strip():
                doc.add_paragraph(line)
        doc.add_heading("Transcript", level=1)
        for line in transcript.splitlines():
            if line.strip():
                doc.add_paragraph(line)
        doc.save(destination)
    elif suffix == ".pdf":
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.pdfgen import canvas
        except ImportError as exc:
            raise RuntimeError("PDF export requires reportlab. Install dependencies from requirements.txt.") from exc
        pdf = canvas.Canvas(str(destination), pagesize=letter)
        font_name = "Helvetica"
        windows = Path(os.environ.get("WINDIR", r"C:\Windows"))
        for font_file in (windows / "Fonts" / "arial.ttf", windows / "Fonts" / "segoeui.ttf"):
            if font_file.exists():
                pdfmetrics.registerFont(TTFont("MeetingNotesUnicode", str(font_file)))
                font_name = "MeetingNotesUnicode"
                break
        pdf.setFont(font_name, 10)
        _, height = letter
        y = height - 48
        for paragraph in content.splitlines():
            for line in _wrap(paragraph, 95):
                if y < 48:
                    pdf.showPage()
                    pdf.setFont(font_name, 10)
                    y = height - 48
                pdf.drawString(48, y, line[:120])
                y -= 14
        pdf.save()
    else:
        raise ValueError("Choose a .txt, .pdf, or .docx file.")
    return destination


def _wrap(text: str, width: int):
    words, line, result = text.split(), "", []
    for word in words:
        if len(line) + len(word) + 1 > width and line:
            result.append(line)
            line = word
        else:
            line = f"{line} {word}".strip()
    result.append(line)
    return result
